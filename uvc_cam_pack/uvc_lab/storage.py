from __future__ import annotations

from dataclasses import asdict
from datetime import datetime, timezone
import hashlib
import importlib.metadata
import json
import math
import os
from pathlib import Path
import platform
import shutil
import sqlite3
import subprocess
import threading
import time
import uuid

import cv2
import numpy as np

from .config import AppConfig, geometry_override


def atomic_json(path: Path, value) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(path.name+".tmp-"+uuid.uuid4().hex)
    try:
        with temp.open("w", encoding="utf-8") as f:
            json.dump(value, f, ensure_ascii=False, indent=2, allow_nan=False)
            f.flush()
            os.fsync(f.fileno())
        os.replace(temp, path)
    finally:
        temp.unlink(missing_ok=True)


def image_bytes(image: np.ndarray, fmt: str, quality: int = 97) -> bytes:
    params = [cv2.IMWRITE_JPEG_QUALITY, quality] if fmt == "jpg" else ([cv2.IMWRITE_PNG_COMPRESSION, 1] if fmt == "png" else [cv2.IMWRITE_TIFF_COMPRESSION, 1])
    ok, data = cv2.imencode("."+fmt, image, params)
    if not ok:
        raise IOError(f"Image encoder failed: {fmt}")
    return data.tobytes()


def write_image(path: Path, image: np.ndarray, quality: int = 97) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    data = image_bytes(image, path.suffix[1:], quality)
    temporary = path.with_suffix(path.suffix+".partial")
    try:
        with temporary.open("wb") as f:
            f.write(data)
            f.flush()
            os.fsync(f.fileno())
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def read_image(path: str | Path) -> np.ndarray:
    data = np.fromfile(str(path), dtype=np.uint8)
    image = cv2.imdecode(data, cv2.IMREAD_COLOR)
    if image is None:
        raise IOError(f"Cannot decode image: {path}")
    return image


def nearest_existing(path: Path) -> Path:
    while not path.exists():
        if path.parent == path:
            raise ValueError("No existing parent for output path")
        path = path.parent
    return path


def estimate_storage(config: AppConfig, free_bytes: int | None = None) -> dict:
    config.validate()
    s, c = config.storage, config.capture
    free = shutil.disk_usage(nearest_existing(Path(s.directory).resolve())).free if free_bytes is None else free_bytes
    total_rate = 0.0
    target_count = len(config.hardware.targets) if config.hardware.enabled else 1
    for cam in config.cameras:
        g = geometry_override(config.geometry, config.geometry.per_camera.get(cam.id, {}))
        w, h = (g.output_size if g.output_size else [cam.width, cam.height])
        area_fraction = sum(r.rect[2]*r.rect[3] for r in g.rois)
        cells = sum(r.rows*r.cols for r in g.rois)
        frame_bytes = (cam.width*cam.height*3 if s.raw else 0)+(w*h*3*area_fraction if s.processed else 0)
        # Conservative image payload; no assumed compression savings for scientific data.
        frame_bytes = frame_bytes*1.05+cells*8192+4096
        if g.stabilize:
            frame_bytes += w*h*area_fraction*1.05
        rate = cam.fps if c.mode == "continuous" else target_count/c.interval_s
        total_rate += frame_bytes*rate
        if s.video:
            video_time_ratio = 1 if c.mode == "continuous" else rate/s.video_fps
            if s.video_container == "avi":
                total_rate += w*h*3*area_fraction*1.1*s.video_fps*video_time_ratio
            else:
                total_rate += s.video_bitrate_mbps*1e6/8*cells*video_time_ratio*1.05
    reserve = c.min_free_gb*1e9
    available = max(0, free-reserve)
    # Include first capture even for a duration shorter than one timelapse interval.
    equivalent_s = c.duration_s
    if c.duration_s is not None and c.mode == "timelapse":
        equivalent_s = math.ceil(c.duration_s/c.interval_s)*c.interval_s
    need = None if equivalent_s is None else total_rate*equivalent_s
    return {"free_bytes": free, "reserve_bytes": reserve, "estimated_bytes_per_second": total_rate,
            "estimated_total_bytes": need, "estimated_hours_available": available/total_rate/3600 if total_rate else None,
            "sufficient": None if need is None else need <= available,
            "unlimited": c.duration_s is None,
            "basis": "Uncompressed BGR8 image payload + 5% + per-cell metadata; capped video bitrate. Analysis caches/exports require extra space."}


def file_hash(path: str | Path) -> str | None:
    p = Path(path)
    if not p.is_file():
        return None
    digest = hashlib.sha256()
    with p.open("rb") as f:
        for block in iter(lambda: f.read(1024*1024), b""):
            digest.update(block)
    return digest.hexdigest()


class RunStore:
    def __init__(self, config: AppConfig, kind: str = "capture"):
        self.config = config
        safe = "".join(c if c.isalnum() or c in "_-" else "_" for c in config.experiment)[:60] or "experiment"
        self.path = Path(config.storage.directory).resolve()/f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{safe}_{uuid.uuid4().hex[:8]}"
        self.path.mkdir(parents=True)
        config.save(self.path/"settings.json", redact=True)
        self.lock = threading.RLock()
        self.closed = False
        self.db = sqlite3.connect(self.path/"records.sqlite", check_same_thread=False)
        self.db.execute("PRAGMA journal_mode=WAL")
        self.db.execute("PRAGMA synchronous=FULL")
        self.db.executescript("""
            CREATE TABLE events(id INTEGER PRIMARY KEY, utc TEXT, kind TEXT, payload TEXT);
            CREATE TABLE frames(id INTEGER PRIMARY KEY, stream TEXT, sequence INTEGER, timestamp REAL,
                monotonic REAL, raw_path TEXT, processed_path TEXT, metadata TEXT, detections TEXT DEFAULT '[]', tracks TEXT DEFAULT '[]');
            CREATE INDEX stream_sequence ON frames(stream, sequence);
            CREATE TABLE sensors(id INTEGER PRIMARY KEY, timestamp REAL, temperature REAL, humidity REAL, raw TEXT);
        """)
        self.manifest = {"schema_version": 1, "kind": kind, "state": "running", "created_utc": datetime.now(timezone.utc).isoformat(),
                         "software": environment(), "raw_definition": "Untransformed decoded UVC BGR8 frame, not sensor Bayer RAW",
                         "timestamps": "Host UTC Unix + monotonic read completion; cameras are not hardware synchronized",
                         "settings": "settings.json", "records": "records.sqlite"}
        self.flush_manifest()
        self.event("run_started", {"kind": kind})

    def flush_manifest(self):
        atomic_json(self.path/"manifest.json", self.manifest)

    def event(self, kind: str, payload: dict):
        with self.lock:
            if self.closed:
                return
            self.db.execute("INSERT INTO events(utc,kind,payload) VALUES(?,?,?)", (datetime.now(timezone.utc).isoformat(), kind, json.dumps(payload, ensure_ascii=False, allow_nan=False)))
            self.db.commit()

    def frame(self, stream: str, sequence: int, timestamp: float, monotonic: float, raw_path: str, processed_path: str, metadata: dict, detections=None, tracks=None):
        with self.lock:
            self.db.execute("INSERT INTO frames(stream,sequence,timestamp,monotonic,raw_path,processed_path,metadata,detections,tracks) VALUES(?,?,?,?,?,?,?,?,?)",
                            (stream, sequence, timestamp, monotonic, raw_path, processed_path, json.dumps(metadata, allow_nan=False), json.dumps(detections or [], allow_nan=False), json.dumps(tracks or [], allow_nan=False)))
            self.db.commit()

    def sensor(self, timestamp, temperature, humidity, raw):
        with self.lock:
            self.db.execute("INSERT INTO sensors(timestamp,temperature,humidity,raw) VALUES(?,?,?,?)", (timestamp, temperature, humidity, raw))
            self.db.commit()

    def finish(self, state: str, summary: dict):
        self.event("run_finished", {"state": state, **summary})
        self.manifest.update(state=state, finished_utc=datetime.now(timezone.utc).isoformat(), summary=summary)
        self.flush_manifest()
        with self.lock:
            self.db.execute("PRAGMA wal_checkpoint(TRUNCATE)")
            self.db.close()
            self.closed = True


def environment() -> dict:
    versions = {}
    for name in ["PyQt6", "opencv-python", "numpy", "scipy", "ultralytics", "torch", "openpyxl", "pyserial"]:
        try:
            versions[name] = importlib.metadata.version(name)
        except importlib.metadata.PackageNotFoundError:
            versions[name] = None
    executable = shutil.which("ffmpeg")
    ffmpeg_version = None
    if executable:
        output = subprocess.run([executable, "-version"], capture_output=True, timeout=5,
                                creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
        ffmpeg_version = output.stdout.decode("utf-8", errors="replace").splitlines()[0]
    return {"python": platform.python_version(), "platform": platform.platform(), "packages": versions,
            "opencv_build": cv2.getBuildInformation(), "ffmpeg": executable, "ffmpeg_version": ffmpeg_version}


class VideoSink:
    """Segmented CFR output. Continuous mode resamples by acquisition timestamps.

    Timelapse mode intentionally writes exactly one video frame per acquired image.
    Counts expose duplicated/dropped frames rather than changing playback silently.
    """
    def __init__(self, base: Path, config: AppConfig, size: tuple[int, int]):
        self.base, self.config, self.size = base, config, size
        self.process = None
        self.part = -1
        self.frames = 0
        self.duplicates = 0
        self.dropped = 0
        self.origin = None
        self.last = None
        self.local_frames = 0
        self.logfile = None
        self.video_path = None

    def _open(self):
        executable = shutil.which("ffmpeg")
        if not executable:
            raise RuntimeError("FFmpeg is required for high-quality MP4/AVI recording")
        self.part += 1
        self.base.parent.mkdir(parents=True, exist_ok=True)
        s = self.config.storage
        self.video_path = self.base.with_name(f"{self.base.name}_{self.part:04d}.{s.video_container}")
        temp = self.video_path.with_name(self.video_path.stem+".partial"+self.video_path.suffix)
        args = [executable, "-hide_banner", "-loglevel", "error", "-nostdin", "-y", "-f", "rawvideo", "-pixel_format", "bgr24",
                "-video_size", f"{self.size[0]}x{self.size[1]}", "-framerate", str(s.video_fps), "-i", "pipe:0", "-an"]
        if s.video_container == "mp4":
            args += ["-vf", "pad=ceil(iw/2)*2:ceil(ih/2)*2", "-c:v", "libx264", "-preset", "veryfast", "-crf", str(s.video_crf),
                     "-maxrate", f"{s.video_bitrate_mbps}M", "-bufsize", f"{s.video_bitrate_mbps*2}M", "-pix_fmt", "yuv420p", "-movflags", "+faststart"]
        else:
            args += ["-c:v", "mjpeg", "-q:v", "2", "-pix_fmt", "yuvj444p"]
        args += [str(temp)]
        self.logfile = self.video_path.with_suffix(".encoder.log").open("wb")
        self.process = subprocess.Popen(args, stdin=subprocess.PIPE, stdout=subprocess.DEVNULL, stderr=self.logfile,
                                        creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
        self.local_frames = 0

    def _write(self, image):
        if self.process is None:
            self._open()
        if image.shape[:2] != (self.size[1], self.size[0]):
            raise ValueError("Video frame size changed")
        self.process.stdin.write(np.ascontiguousarray(image).tobytes())
        self.frames += 1
        self.local_frames += 1
        if self.local_frames >= max(1, round(self.config.storage.segment_s*self.config.storage.video_fps)):
            self._close_segment()

    def write(self, image: np.ndarray, timestamp: float):
        if self.origin is None:
            self.origin = timestamp
        if self.config.capture.mode == "timelapse":
            self._write(image)
        else:
            index = int((timestamp-self.origin)*self.config.storage.video_fps+1e-6)
            if index < self.frames:
                self.dropped += 1
            else:
                # Limit damage from a clock/driver discontinuity; retain evidence.
                if index-self.frames > self.config.storage.video_fps*10:
                    raise RuntimeError("Camera stalled >10s; refusing to fabricate a long video gap")
                while self.frames < index and self.last is not None:
                    self._write(self.last)
                    self.duplicates += 1
                self._write(image)
        self.last = image.copy()

    def _close_segment(self):
        if self.process is None:
            return
        process = self.process
        self.process = None
        try:
            process.stdin.close()
            code = process.wait(timeout=30)
        except BaseException:
            process.kill()
            process.wait()
            raise
        finally:
            self.logfile.close()
        if code:
            raise RuntimeError(f"FFmpeg encoder failed ({code}); see {self.video_path.with_suffix('.encoder.log')}")
        temp = self.video_path.with_name(self.video_path.stem+".partial"+self.video_path.suffix)
        os.replace(temp, self.video_path)

    def close(self):
        self._close_segment()
        return {"frames": self.frames, "duplicates": self.duplicates, "dropped_by_resampler": self.dropped, "segments": self.part+1}


def export_excel(run_path: Path) -> Path:
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    db = sqlite3.connect(run_path/"records.sqlite")
    book = Workbook(write_only=True)
    frames = book.create_sheet("Frames")
    frames.append(["stream", "sequence", "unix_s", "raw", "processed", "metadata_json"])
    detections = book.create_sheet("Detections")
    detections.append(["stream", "sequence", "unix_s", "object", "class", "confidence", "x1_px", "y1_px", "x2_px", "y2_px", "area_px2", "area_um2", "payload_json"])
    tracks = book.create_sheet("Tracks")
    tracks.append(["stream", "sequence", "unix_s", "track_id", "payload_json"])
    row_counts = {"Detections": 1, "Frames": 1, "Tracks": 1}
    sheets = {"Detections": detections, "Frames": frames, "Tracks": tracks}
    parts = {key: 1 for key in sheets}
    headers = {
        "Frames": ["stream", "sequence", "unix_s", "raw", "processed", "metadata_json"],
        "Detections": ["stream", "sequence", "unix_s", "object", "class", "confidence", "x1_px", "y1_px", "x2_px", "y2_px", "area_px2", "area_um2", "payload_json"],
        "Tracks": ["stream", "sequence", "unix_s", "track_id", "payload_json"],
        "Environment": ["unix_s", "temperature_C", "humidity_percent", "source_line"],
    }
    def sensor_chart(sheet, rows):
        if rows <= 2:
            return
        chart = LineChart()
        chart.title = "Temperature / humidity"
        chart.add_data(Reference(sheet, min_col=2, max_col=3, min_row=1, max_row=rows), titles_from_data=True)
        chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=rows))
        sheet.add_chart(chart, "F2")
    def append(key, row):
        if row_counts[key] >= 1_048_576:
            if key == "Environment":
                sensor_chart(sheets[key], row_counts[key])
            parts[key] += 1
            sheets[key] = book.create_sheet(f"{key}_{parts[key]}")
            sheets[key].append(headers[key])
            row_counts[key] = 1
        # Prevent strings from user notes/labels becoming spreadsheet formulas.
        sheets[key].append(["'"+x if isinstance(x, str) and x.startswith(("=", "+", "-", "@")) else x for x in row])
        row_counts[key] += 1
    try:
        for stream, seq, ts, raw, processed, meta, ds, ts_data in db.execute("SELECT stream,sequence,timestamp,raw_path,processed_path,metadata,detections,tracks FROM frames ORDER BY id"):
            append("Frames", [stream, seq, ts, raw, processed, meta])
            for idx, det in enumerate(json.loads(ds)):
                append("Detections", [stream, seq, ts, idx, det["label"], det["confidence"], *det["bbox"], det.get("area_px2"), det.get("area_um2"), json.dumps(det, ensure_ascii=False)[:32767]])
            for track in json.loads(ts_data):
                append("Tracks", [stream, seq, ts, track["id"], json.dumps(track, ensure_ascii=False)[:32767]])
        # Write-only rows bound memory even for multi-day experiments.
        sensor = book.create_sheet("Environment")
        sensor.append(headers["Environment"])
        sheets["Environment"], row_counts["Environment"], parts["Environment"] = sensor, 1, 1
        for row in db.execute("SELECT timestamp,temperature,humidity,raw FROM sensors ORDER BY id"):
            append("Environment", list(row))
        sensor_chart(sheets["Environment"], row_counts["Environment"])
        dest = run_path/"report.xlsx"
        book.save(dest)
        return dest
    finally:
        db.close()
